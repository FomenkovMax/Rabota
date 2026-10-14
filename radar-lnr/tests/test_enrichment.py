import io

import httpx
import openpyxl
import pytest

from radar.bot.inn_file import MAX_INNS, parse_inn_file
from radar.db import DB
from radar.enrichment.base import AdapterError, LookupResult, OrgCard
from radar.enrichment.dadata import DaDataAdapter, parse_party
from radar.enrichment.service import Enricher
from radar.report.excel import cards_workbook

SBER = "7707083893"
GAZPROM = "7736050003"

PARTY = {"value": 'ООО "Ромашка"', "data": {
    "inn": SBER, "ogrn": "1027700132195", "kpp": "773601001", "okved": "64.19", "employee_count": 250,
    "name": {"full_with_opf": 'Общество с ограниченной ответственностью "Ромашка"', "short_with_opf": 'ООО "Ромашка"'},
    "address": {"value": "г Луганск, ул Ленина, д 1"}, "management": {"name": "Иванов Иван Иванович", "post": "Директор"},
    "state": {"status": "ACTIVE"}, "founders": [{"name": "ООО Холдинг"}],
    "phones": [{"value": "+7 642 000-00-00"}], "emails": [{"value": "info@romashka.example"}], "branch_type": "MAIN"}}


def adapter(handler):
    return DaDataAdapter("k", httpx.AsyncClient(transport=httpx.MockTransport(handler)))


def test_parse_party_maps_fields_and_leaves_missing_empty():
    c = parse_party(PARTY)
    assert c.inn == SBER and c.director == "Иванов Иван Иванович" and c.employees == 250 and c.status == "Действует"
    bare = parse_party({"value": "X", "data": {"inn": GAZPROM}})
    assert bare.director is None and bare.employees is None and bare.founders == []   # ничего не выдумываем


async def test_lookup_inn_found_and_cached(tmp_path):
    calls = []

    def handler(req):
        calls.append(1)
        return httpx.Response(200, json={"suggestions": [PARTY]})

    enr = Enricher(DB(tmp_path / "t.db"), [adapter(handler)])
    r1 = await enr.lookup_inn(SBER)
    r2 = await enr.lookup_inn(SBER)
    assert r1.status == "found" and r1.card.name_short == 'ООО "Ромашка"' and r2.status == "found"
    assert len(calls) == 1                         # второй раз — из кэша


async def test_bad_inn_checksum_rejected_without_api_call(tmp_path):
    enr = Enricher(DB(tmp_path / "t.db"), [adapter(lambda r: pytest.fail("не должно быть запроса"))])
    assert (await enr.lookup_inn("7707083894")).status == "error"


async def test_not_found_and_api_errors(tmp_path):
    enr = Enricher(DB(tmp_path / "t.db"), [adapter(lambda r: httpx.Response(200, json={"suggestions": []}))])
    assert (await enr.lookup_inn(SBER)).status == "not_found"
    bad = Enricher(DB(tmp_path / "t2.db"), [adapter(lambda r: httpx.Response(403))])
    res = await bad.lookup_inn(SBER)
    assert res.status == "error" and "DADATA_API_KEY" in res.message


async def test_name_multiple_candidates_needs_review_top3(tmp_path):
    items = [{"value": f"ООО Ромашка {i}", "data": {"inn": f"77070838{i}0", "name": {"full_with_opf": f"ООО Ромашка {i}"}}} for i in range(5)]
    enr = Enricher(DB(tmp_path / "t.db"), [adapter(lambda r: httpx.Response(200, json={"suggestions": items}))])
    res = await enr.lookup_name("Ромашка")
    assert res.status == "needs_review" and len(res.candidates) == 3 and res.card is None   # не угадываем


async def test_name_single_exact_candidate_found(tmp_path):
    enr = Enricher(DB(tmp_path / "t.db"), [adapter(lambda r: httpx.Response(200, json={"suggestions": [PARTY]}))])
    assert (await enr.lookup_name('ООО "Ромашка"')).status == "found"


async def test_no_adapters_is_disabled(tmp_path):
    assert (await Enricher(DB(tmp_path / "t.db"), []).lookup_inn(SBER)).status == "disabled"


def test_inn_file_txt_csv_dedup_and_invalid():
    raw = f"ИНН\n{SBER}\n{GAZPROM};{SBER}\n1234567890\nмусор 12\n".encode("utf-8")
    valid, invalid, trunc = parse_inn_file("list.csv", raw)
    assert valid == [SBER, GAZPROM] and invalid == ["1234567890"] and not trunc


def test_inn_file_xlsx_numeric_cells_and_limit():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ИНН"])
    ws.append([int(SBER)])                 # число, а не строка
    buf = io.BytesIO()
    wb.save(buf)
    assert parse_inn_file("a.xlsx", buf.getvalue())[0] == [SBER]
    with pytest.raises(ValueError):
        parse_inn_file("a.pdf", b"")


def test_inn_limit_50():
    # генерируем 60 валидных ИНН подбором контрольной цифры
    from radar.llm.verify import inn_checksum_ok
    out = []
    n = 770000000
    while len(out) < 60:
        for d in range(10):
            c = f"{n}{d}"
            if inn_checksum_ok(c):
                out.append(c)
        n += 1
    valid, _, trunc = parse_inn_file("a.txt", "\n".join(out).encode())
    assert len(valid) == MAX_INNS and trunc


def test_excel_cards_workbook_has_nd_and_review_sheet():
    cand = OrgCard(inn=SBER, name_full="ООО Ромашка 1")
    data = cards_workbook([(SBER, LookupResult("found", card=parse_party(PARTY))),
                           ("Ромашка", LookupResult("needs_review", candidates=[cand], message="Требует проверки")),
                           ("1234567890", LookupResult("error", message="неверный ИНН"))])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    rows = list(wb["Карточки"].iter_rows(values_only=True))
    assert rows[1][1] == "Найдено" and rows[1][7] == "Иванов Иван Иванович"
    assert rows[2][1] == "Требует проверки" and rows[2][2] == "н/д"
    assert wb["Требует проверки (топ-3)"].max_row == 2


async def test_name_search_is_restricted_to_lnr_by_kladr(tmp_path):
    seen = []

    def handler(req):
        import json as j
        seen.append(j.loads(req.content))
        return httpx.Response(200, json={"suggestions": []})

    res = await Enricher(DB(tmp_path / "t.db"), [adapter(handler)]).lookup_name("Клеверфарм")
    assert seen[0]["locations"] == [{"kladr_id": "9400000000000"}]
    assert res.status == "not_found"      # нет в ЛНР — не подставляем московских однофамильцев


# ---------- v3: сравнение названий, подсказки вне ЛНР, несколько организаций ----------
def _party(name, inn, addr="Луганская Народная Респ, г Луганск"):
    return {"value": name, "data": {"inn": inn, "name": {"full_with_opf": name, "short_with_opf": name},
                                    "address": {"value": addr}}}


def _by_scope(lnr_items, wide_items):
    """Заглушка DaData: с фильтром kladr_id отдаёт ЛНР-список, без фильтра — общий."""
    def handler(req):
        import json as j
        body = j.loads(req.content)
        items = lnr_items if body.get("locations") else wide_items
        return httpx.Response(200, json={"suggestions": items})
    return handler


async def test_subset_name_is_not_a_match_political_group_vs_company(tmp_path):
    # «ПГ Родина» ⊂ «Промышленная группа Родина»: раньше это давало ложное совпадение
    wide = [_party('АО "ПРОМЫШЛЕННАЯ ГРУППА "РОДИНА"', "7733358632", "г Москва, ул Молдавская")]
    enr = Enricher(DB(tmp_path / "t.db"), [adapter(_by_scope([], wide))])
    res = await enr.lookup_name("ПГ «Родина»")
    assert res.status == "needs_review"          # только подсказка, не «found»
    assert res.card is None and "вне ЛНР" in res.message


async def test_outside_lnr_is_only_a_hint_even_for_exact_name(tmp_path):
    wide = [_party('ООО "ЭНЕРГОСБЫТ ЛУГАНСК"', "9725129223", "г Москва, проезд Проектируемый")]
    res = await Enricher(DB(tmp_path / "t.db"), [adapter(_by_scope([], wide))]).lookup_name('ООО "Энергосбыт Луганск"')
    assert res.status == "needs_review" and res.candidates[0].inn == "9725129223"


async def test_clear_winner_in_lnr_is_found(tmp_path):
    lnr = [_party('ГУП ЛНР "Лугансктеплокоммунэнерго"', "9400000001"),
           _party("ПЕРВИЧНАЯ ПРОФСОЮЗНАЯ ОРГАНИЗАЦИЯ РАБОТНИКОВ ЛУГАНСКТЕПЛОКОММУНЭНЕРГО", "9400000002")]
    res = await Enricher(DB(tmp_path / "t.db"), [adapter(_by_scope(lnr, []))]).lookup_name("Лугансктеплокоммунэнерго")
    assert res.status == "found" and res.card.inn == "9400000001"


async def test_ambiguous_in_lnr_shows_top3_with_inns(tmp_path):
    lnr = [_party(f'ГБОУ ЛНР "Краснолучская школа № {i}"', f"94100000{i}") for i in (3, 4, 5, 6)]
    res = await Enricher(DB(tmp_path / "t.db"), [adapter(_by_scope(lnr, []))]).lookup_name("Краснолучская школа № 7")
    assert res.status == "needs_review" and len(res.candidates) == 3


async def test_enrich_event_prefers_found_org_among_first_three(tmp_path):
    lnr = [_party('ООО "Ромашка"', SBER)]

    def handler(req):
        import json as j
        q = j.loads(req.content)["query"]
        return httpx.Response(200, json={"suggestions": lnr if "Ромашка" in q else []})

    enr = Enricher(DB(tmp_path / "t.db"), [adapter(handler)])
    ev = {"organizations": [{"name": "Неизвестная контора"}, {"name": 'ООО "Ромашка"'}]}
    await enr.enrich_event(ev)
    assert ev["enriched"]["status"] == "found" and ev["enriched"]["matched_org"] == 'ООО "Ромашка"'


# ---------------- v4: скобки в названии, профсоюзы и головная организация ----------------
async def test_brackets_are_stripped_from_search_query(tmp_path):
    queries = []

    def handler(req):
        import json as j
        queries.append(j.loads(req.content)["query"])
        return httpx.Response(200, json={"suggestions": [_party('ООО "Луганскгаз"', "9704210635")]})

    await Enricher(DB(tmp_path / "t.db"), [adapter(handler)]).lookup_name("Луганскгаз (филиал ООО «ЧМНГ»)")
    assert queries[0] == "Луганскгаз"                  # без скобок; исходный запрос — только если пусто


async def test_union_orgs_go_last_and_head_org_goes_first(tmp_path):
    cands = [_party('ППО "ЛУГАНСКГАЗ" ФИЛИАЛ ООО "ЧМНГ"', "9402003047"),
             _party('ЦПО УМГ "ЛУГАНСКГАЗ" ФИЛИАЛ ООО "ЧМНГ"', "9403005713"),
             _party('"ЛУГАНСКГАЗ" ФИЛИАЛ ОБЩЕСТВА С ОГРАНИЧЕННОЙ ОТВЕТСТВЕННОСТЬЮ "ЧЕРНОМОРНЕФТЕГАЗ"', "9704210635")]
    cands[0]["data"]["name"] = {"full_with_opf": "ПЕРВИЧНАЯ ПРОФСОЮЗНАЯ ОРГАНИЗАЦИЯ ЛУГАНСКГАЗ", "short_with_opf": 'ППО "ЛУГАНСКГАЗ"'}
    cands[1]["data"]["name"] = {"full_with_opf": "ЦЕХОВАЯ ПРОФСОЮЗНАЯ ОРГАНИЗАЦИЯ ЛУГАНСКГАЗ", "short_with_opf": 'ЦПО УМГ "ЛУГАНСКГАЗ"'}
    cands[2]["data"]["name"] = {"full_with_opf": 'ФИЛИАЛ ОБЩЕСТВА С ОГРАНИЧЕННОЙ ОТВЕТСТВЕННОСТЬЮ "ЧЕРНОМОРНЕФТЕГАЗ" "ЛУГАНСКГАЗ"'}
    res = await Enricher(DB(tmp_path / "t.db"), [adapter(_by_scope(cands, []))]).lookup_name("Луганскгаз")
    # единственный работодатель (профсоюзные ППО/ЦПО не считаются) -> принимаем головную организацию
    assert res.status == "found" and res.card.inn == "9704210635"


async def test_head_org_and_branches_with_same_inn_count_as_one(tmp_path):
    head = _party('ГУП ЛНР "ЛУГАНСКВОДА"', "9403001349")
    head["data"]["name"] = {"full_with_opf": 'ГОСУДАРСТВЕННОЕ УНИТАРНОЕ ПРЕДПРИЯТИЕ ЛНР "ЛУГАНСКВОДА"', "short_with_opf": 'ГУП ЛНР "ЛУГАНСКВОДА"'}
    branches = []
    for city in ("АЛЧЕВСКОЕ УПРАВЛЕНИЕ", "АНТРАЦИТОВСКИЙ ДЕПАРТАМЕНТ"):
        b = _party(f"{city} ГУП", "9403001349")
        b["data"]["name"] = {"full_with_opf": f'{city} ГУП ЛНР "ЛУГАНСКВОДА"'}
        branches.append(b)
    res = await Enricher(DB(tmp_path / "t.db"), [adapter(_by_scope([head, *branches], []))]).lookup_name("Лугансквода")
    assert res.status == "found" and res.card.inn == "9403001349"        # один ИНН — это одна организация


async def test_union_org_does_not_block_single_employer(tmp_path):
    head = _party('ГУП ЛНР "ЛУГАНСКВОДА"', "9403001349")
    head["data"]["name"] = {"full_with_opf": 'ГОСУДАРСТВЕННОЕ УНИТАРНОЕ ПРЕДПРИЯТИЕ ЛНР "ЛУГАНСКВОДА"', "short_with_opf": 'ГУП ЛНР "ЛУГАНСКВОДА"'}
    union = _party('ОППО "ЛУГАНСКВОДА"', "9403015140")
    union["data"]["name"] = {"full_with_opf": 'ОБЪЕДИНЕННАЯ ПЕРВИЧНАЯ ПРОФСОЮЗНАЯ ОРГАНИЗАЦИЯ "ЛУГАНСКВОДА"', "short_with_opf": 'ОППО "ЛУГАНСКВОДА"'}
    res = await Enricher(DB(tmp_path / "t.db"), [adapter(_by_scope([union, head], []))]).lookup_name("Лугансквода")
    assert res.status == "found" and res.card.inn == "9403001349"
