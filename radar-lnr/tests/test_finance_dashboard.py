import io
import json
import re
from datetime import datetime, timedelta, timezone

import openpyxl

from radar.finance import effect_text, estimate
from radar.report.dashboard import render_dashboard
from radar.report.excel import events_workbook

SAL = {"default": 40000, "by_okved2": {"64": 100000}}
FIN = {"salary_share": 0.5, "package_yield": 0.02, "months": 12}


def ev(**kw):
    base = {"category": 1, "title": "T", "summary": "S", "relevance_score": 4, "organizations": [{"name": "ООО А"}],
            "persons": [], "source_urls": ["https://t.me/a/1"], "channels": ["a"]}
    return {**base, **kw}


def test_finance_formula_with_okved():
    e = ev(employees_mentioned=100, enriched={"status": "found", "card": {"okved": "64.19", "employees": None}})
    f = estimate(e, SAL, FIN)
    assert f["fot_month"] == 100 * 100000 and f["fot_year"] == 120_000_000
    assert f["potential_year"] == 120_000_000 * 0.5 * 0.02 and f["employees_source"] == "текст новости"
    assert "≈" in effect_text(f) and "оценка" in effect_text(f)


def test_finance_uses_default_salary_and_prefers_dadata_employees():
    f = estimate(ev(employees_mentioned=5, enriched={"status": "found", "card": {"okved": None, "employees": 200}}), SAL, FIN)
    assert f["employees"] == 200 and f["employees_source"] == "DaData" and f["salary"] == 40000
    assert "заглушка" in f["salary_source"]


def test_finance_nd_without_employees():
    f = estimate(ev(), SAL, FIN)
    assert f["potential_year"] is None and f["fot_year"] is None
    assert effect_text(f).startswith("н/д")


def test_dashboard_embeds_safe_data_and_nd():
    evil = ev(title="</script><img src=x onerror=alert(1)>", fin_effect_text="н/д", finance=estimate(ev(), SAL, FIN))
    now = datetime.now(timezone.utc)
    html = render_dashboard([evil], now - timedelta(hours=4), now, {"a": "Луганск"})
    assert html.count("</script>") == 2                    # только наши два закрывающих тега
    blob = re.search(r'id="data" type="application/json">(.*?)</script>', html, re.S).group(1)
    data = json.loads(blob)
    assert data["events"][0]["title"].startswith("</script>") and data["events"][0]["district"] == "Луганск"
    assert "innerHTML" not in html                        # данные только через textContent


def test_events_excel_has_formula_sheet_and_nd():
    data = events_workbook([ev(finance=estimate(ev(), SAL, FIN))], {"a": "Луганск"})
    wb = openpyxl.load_workbook(io.BytesIO(data))
    row = list(wb["События"].iter_rows(values_only=True))[1]
    assert row[6] == "Луганск" and row[8] == "н/д" and row[10] == "н/д"
    assert "Формула" in wb.sheetnames
