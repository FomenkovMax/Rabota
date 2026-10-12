"""Excel-выгрузки: карточки организаций (/inn, пакетная проверка)."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..enrichment.base import LookupResult

STATUS_RU = {"found": "Найдено", "needs_review": "Требует проверки", "not_found": "Не найдено",
             "error": "Ошибка", "disabled": "Обогащение выключено"}
COLS = [("Запрос", None), ("Статус", None), ("ИНН", "inn"), ("ОГРН", "ogrn"), ("Наименование", "name_full"),
        ("ОКВЭД", "okved"), ("Адрес", "address"), ("Руководитель (ЛПР)", "director"), ("Должность", "director_post"),
        ("Учредители", "founders"), ("Числ. (среднесписочная)", "employees"), ("Статус орг.", "status"),
        ("Телефоны", "phones"), ("Email", "emails"), ("Примечание", None)]
NA = "н/д"


def _val(v):
    if v in (None, "", []):
        return NA
    return "; ".join(v) if isinstance(v, list) else v


def _style(ws, widths):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="21A038")     # зелёный Сбера
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def cards_workbook(items: list[tuple[str, LookupResult]]) -> bytes:
    """items: [(исходный запрос, результат)] -> .xlsx в байтах. Лист 2 — кандидаты для «Требует проверки»."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Карточки"
    ws.append([c[0] for c in COLS])
    review = []
    for query, res in items:
        card = res.card.to_dict() if res.card else {}
        row = [query, STATUS_RU.get(res.status, res.status)]
        row += [_val(card.get(k)) if k else "" for _, k in COLS[2:-1]]
        row.append(res.message or "")
        ws.append(row)
        for i, cand in enumerate(res.candidates, 1):
            review.append([query, i, cand.inn or NA, cand.name_full or NA, cand.address or NA, cand.director or NA])
    _style(ws, [16, 18, 14, 16, 40, 10, 40, 28, 22, 30, 14, 14, 22, 26, 36])
    if review:
        w2 = wb.create_sheet("Требует проверки (топ-3)")
        w2.append(["Запрос", "№", "ИНН", "Наименование", "Адрес", "Руководитель"])
        for r in review:
            w2.append(r)
        _style(w2, [28, 5, 14, 44, 44, 28])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
