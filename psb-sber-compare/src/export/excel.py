"""Свод в Excel.

Листы:
  Светофор   — пары продуктов, дельты, вердикт;
  Акции      — действующие предложения по сегментам;
  Изменения  — что поменялось с прошлого сбора;
  Продукты   — полная выгрузка по каждому банку.

Формат выбран под то, как им будут пользоваться: цифры дособирают руками
и пересылают дальше, поэтому колонки числовые, а не текстовые, и на шапке
стоит закреплённая строка с автофильтром.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..compare import GREEN, GREY, LIGHT_LABEL, RED, YELLOW

log = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)

LIGHT_FILL = {
    RED: PatternFill("solid", fgColor="FBD5D5"),
    YELLOW: PatternFill("solid", fgColor="FDF1CD"),
    GREEN: PatternFill("solid", fgColor="D7F0D7"),
    GREY: PatternFill("solid", fgColor="EFEFEF"),
}

THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _write_header(sheet: Any, titles: list[str], widths: list[int]) -> None:
    sheet.append(titles)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 30


def _finish(sheet: Any, columns: int) -> None:
    if sheet.max_row > 1:
        sheet.auto_filter.ref = f"A1:{get_column_letter(columns)}{sheet.max_row}"
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _sheet_traffic(book: Workbook, data: dict[str, Any]) -> None:
    sheet = book.active
    sheet.title = "Светофор"
    _write_header(
        sheet,
        ["Продукт", "Категория", "Банк-конкурент", "Ставка конкурента",
         "Ставка Сбера", "Дельта, п.п.", "Вердикт", "Комментарий"],
        [34, 20, 14, 17, 15, 13, 16, 46],
    )

    for c in data["comparisons"]:
        row = [
            c.label,
            c.category,
            getattr(c.psb, "bank", "") if c.psb else "",
            c.psb_rate,
            c.sber_rate,
            c.delta_rate,
            LIGHT_LABEL[c.light],
            c.reason,
        ]
        sheet.append(row)
        fill = LIGHT_FILL[c.light]
        for cell in sheet[sheet.max_row]:
            cell.fill = fill
        # Числовые колонки — числами, иначе в Excel по ним не посчитать.
        for column in (4, 5, 6):
            cell = sheet.cell(row=sheet.max_row, column=column)
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.00"

    _finish(sheet, 8)


def _sheet_promos(book: Workbook, data: dict[str, Any]) -> None:
    sheet = book.create_sheet("Акции")
    _write_header(
        sheet,
        ["Сегмент", "Вердикт", "Вывод", "Банк", "Предложение",
         "Выгода", "Тип выгоды", "Действует до"],
        [26, 15, 52, 10, 50, 14, 14, 15],
    )

    verdict_light = {"red": RED, "yellow": YELLOW, "green": GREEN, "grey": GREY}
    for segment in data["segments"]:
        fill = LIGHT_FILL[verdict_light.get(segment.verdict, GREY)]
        promos = list(segment.psb) + list(segment.sber)
        if not promos:
            sheet.append([segment.segment, segment.verdict, segment.headline,
                          "", "Предложений нет", "", "", ""])
            for cell in sheet[sheet.max_row]:
                cell.fill = fill
            continue
        for promo in promos:
            sheet.append([
                segment.segment, segment.verdict, segment.headline,
                promo.bank, promo.title, promo.benefit_display,
                promo.benefit_label, promo.valid_display,
            ])
            for cell in sheet[sheet.max_row]:
                cell.fill = fill

    _finish(sheet, 8)


def _sheet_changes(book: Workbook, data: dict[str, Any]) -> None:
    from ..changes import KIND_LABEL

    sheet = book.create_sheet("Изменения")
    _write_header(
        sheet,
        ["Банк", "Что произошло", "Продукт", "Поле", "Было", "Стало",
         "Дельта, п.п.", "Важность"],
        [10, 22, 40, 22, 34, 34, 13, 12],
    )

    for change in data["changes"]:
        sheet.append([
            change["bank"],
            KIND_LABEL.get(change["kind"], change["kind"]),
            change.get("title", ""),
            change.get("field", ""),
            change.get("old_value", ""),
            change.get("new_value", ""),
            change.get("delta"),
            change.get("severity", ""),
        ])
        if change.get("severity") == "high":
            for cell in sheet[sheet.max_row]:
                cell.fill = LIGHT_FILL[RED]

    _finish(sheet, 8)


def _sheet_products(book: Workbook, data: dict[str, Any]) -> None:
    sheet = book.create_sheet("Продукты")
    _write_header(
        sheet,
        ["Банк", "Категория", "Продукт", "Ставка от", "Ставка до",
         "Условия лучшей ставки", "ПСК от", "Сумма до", "Срок до, мес",
         "Источник"],
        [10, 22, 40, 12, 12, 40, 11, 14, 14, 44],
    )

    for product in data["products"]:
        sheet.append([
            product.bank,
            product.category,
            product.title,
            product.rate_min,
            product.rate_max,
            getattr(product, "rate_conditions", ""),
            product.apr_min,
            product.amount_max,
            product.term_max_months,
            product.source_url,
        ])
        for column in (4, 5, 7):
            cell = sheet.cell(row=sheet.max_row, column=column)
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.00"
        cell = sheet.cell(row=sheet.max_row, column=8)
        if isinstance(cell.value, (int, float)):
            cell.number_format = "# ##0"

    _finish(sheet, 10)


def build(path: str | Path, data: dict[str, Any]) -> Path:
    """Собирает книгу Excel со сводом."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    book = Workbook()
    _sheet_traffic(book, data)
    _sheet_promos(book, data)
    _sheet_changes(book, data)
    _sheet_products(book, data)

    book.properties.title = f"Сравнение розничных продуктов — {data['region']}"
    book.save(path)
    log.info("Excel собран: %s", path)
    return path
