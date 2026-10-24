"""Данные с чужих сайтов считаем недоверенными: названия продуктов не должны исполняться ни в HTML, ни в Telegram, ни в Excel."""

import io
import json
import re
from decimal import Decimal as D

import pytest
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from test_reports import NOW  # noqa: F401

from bankintel.analytics.compare import Scenario, compare
from bankintel.analytics.facts import facts_from_params
from bankintel.bot import compare_views
from bankintel.reports import charts
from bankintel.reports.render_html import _render, render_html
from bankintel.reports.render_xlsx import render_xlsx

EVIL = ['</script><img src=x onerror=alert(1)>', '"><svg onload=alert(1)>', "=HYPERLINK(\"http://evil.test\",\"клик\")", "+cmd|' /C calc'!A0", "@SUM(1+1)"]


def _facts(name):
    card = {"title": f"Вклад «{name}»", "headline_rate_pct": {"min": "10", "max": "12"}, "headline_rate_text": "От 10% до 12%",
            "term_text": "91–367 дней", "amount_text": "От 10 000 ₽", "currency": "RUB", "description": name, "product_url": "https://www.psbank.ru/x"}
    f = facts_from_params(bank={"slug": "psb", "name": "ПСБ"}, product={"id": 1, "name": name, "category": "deposit"}, card=card)
    f.document_ids, f.snapshot_ids, f.checked_at = [1], [1], NOW
    return f


def _dataset(name):
    from bankintel.reports.dataset import LIMITATIONS, METHODOLOGY, _product, _scenario_block

    facts = [_facts(name)]
    return {
        "meta": {"title": "t", "region": "ЛНР", "app_version": "x", "method_version": "1", "data_from": NOW.isoformat(), "data_to": NOW.isoformat(),
                 "products_total": 1, "lnr_confirmed": 0, "banks_with_data": 1, "banks_total": 1, "dataset_id": "abc", "generated_at": NOW.isoformat()},
        "coverage": [{"slug": "psb", "name": name, "is_benchmark": False, "status": "data", "deposits": 1, "savings": 0, "manual_documents": 0,
                      "domain_verified": False, "lnr_presence_verified": False, "reason": name, "last_run": None, "last_fetch": None}],
        "scenarios": [_scenario_block(facts, Scenario(D(100000), 6), [("psb", "ПСБ")], NOW)], "insights": [name],
        "changes": {"window_days": 30, "monitoring": {"since": NOW.isoformat(), "last_run": None, "comparisons": 1, "has_data": True},
                    "items": [{"detected_at": NOW.isoformat(), "bank": name, "product": name, "type": "changed", "severity": "rate", "text": name}]},
        "catalog": {"deposit": [_product(facts[0])], "savings": []},
        "sources": [{"document_id": 1, "bank": name, "source_key": "k", "url": "https://www.psbank.ru/x?a=<b>", "origin": "fetched",
                     "fetched_at": NOW.isoformat(), "sha256": "a" * 64, "size_bytes": 1}],
        "methodology": METHODOLOGY, "limitations": LIMITATIONS,
    }


def _scripts(html):
    return len(BeautifulSoup(html, "html.parser").find_all("script"))


@pytest.mark.parametrize("name", EVIL[:2])
def test_html_report_does_not_execute_untrusted_names(name):
    html = render_html(_dataset(name))
    baseline = render_html(_dataset("Обычный вклад"))
    assert _scripts(html) == _scripts(baseline)                        # лишних <script> не появилось
    soup = BeautifulSoup(html, "html.parser")
    assert not soup.find_all("img", onerror=True) and not soup.find_all(onload=True)
    assert not [t for t in soup.find_all(True) if any(a.lower().startswith("on") for a in t.attrs)]  # ни одного обработчика событий


@pytest.mark.parametrize("name", EVIL[:2])
def test_svg_chart_escapes_labels(name):
    ds = _dataset(name)
    svg = charts.svg_range_chart(charts.chart_rows(ds["scenarios"][0]))
    assert svg and "&lt;" in svg and "<img" not in svg   # опасные символы превращены в сущности
    soup = BeautifulSoup(svg, "html.parser")
    assert not soup.find_all("img") and not [t for t in soup.find_all(True) if any(a.startswith("on") for a in t.attrs)]
    assert _scripts(_render(ds, "pdf", lambda rows, uid: charts.svg_range_chart(rows, uid=uid))) == 0  # в PDF-версии скриптов нет вовсе


@pytest.mark.parametrize("name", EVIL)
def test_telegram_texts_escape_html(name):
    facts = [_facts(name)]
    for text in (compare_views.deposits_text(facts), compare_views.compare_text(compare(facts, Scenario(D(100000), 6), now=NOW))):
        assert "<img" not in text and "<svg" not in text and "</script>" not in text
        stripped = re.sub(r"</?(b|i|code)>", "", text)               # остаются только разрешённые Telegram-теги
        assert "<" not in stripped.replace("&lt;", "")


@pytest.mark.parametrize("name", EVIL[2:])
def test_xlsx_formula_injection_neutralised(name):
    wb = load_workbook(io.BytesIO(render_xlsx(_dataset(name))))
    formulas = [(ws.title, c.coordinate, c.value) for ws in wb for row in ws.iter_rows() for c in row if c.data_type == "f"]
    assert formulas == [], f"ячейки-формулы из недоверенных данных: {formulas}"
    texts = [c.value for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any(name in t for t in texts)  # текст сохранён как текст


def test_dataset_json_roundtrip_is_plain_data():
    json.dumps(_dataset(EVIL[0]), ensure_ascii=False)


def test_pdf_renderer_never_loads_external_resources():
    """Модель отказа экранирования: в HTML попали <img src=file://…> и внешний CSS. Рендерер не должен их загружать."""
    from bankintel.reports import render_html

    deny = render_html._make_deny_fetcher()
    calls = []
    original = deny.fetch

    def spy(self, url, headers=None):
        calls.append(url)
        return original(self, url, headers)

    deny.fetch = spy
    from weasyprint import HTML

    html = '<html><head><link rel="stylesheet" href="http://evil.test/x.css"></head><body><img src="file:///etc/passwd"><p>текст</p></body></html>'
    pdf = HTML(string=html, url_fetcher=deny()).write_pdf()
    assert pdf.startswith(b"%PDF")  # документ всё равно строится
    assert any("file:///etc/passwd" in u for u in calls) and any("evil.test" in u for u in calls)  # попытки были, но отклонены


def test_real_report_pdf_makes_no_external_fetch_attempts(monkeypatch):
    from bankintel.reports import render_html

    deny = render_html._make_deny_fetcher()
    monkeypatch.setattr(render_html, "_make_deny_fetcher", lambda: deny)
    deny.attempted.clear()
    render_html.render_pdf(_dataset("Обычный вклад"))
    assert deny.attempted == []  # штатный отчёт не обращается к внешним ресурсам вовсе
