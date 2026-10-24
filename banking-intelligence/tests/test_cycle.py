from datetime import UTC, datetime
from unittest.mock import AsyncMock, patch

import pytest
from aiogram.filters import CommandObject
from sqlalchemy import select

from bankintel.bot import handlers
from bankintel.bot.changes_views import changes_text, parse_days
from bankintel.config import Settings
from bankintel.cycle import format_events, redis_lock, run_cycle, telegram_notifier
from bankintel.db.models import Bank
from bankintel.monitor import monitoring_status
from bankintel.reports.dataset import build_dataset

from test_monitor import DEPOSITS, Pages, mutate


def settings(tmp_path, **kw):
    return Settings(_env_file=None, raw_storage_dir=tmp_path, **kw)


class Fetch(Pages):
    async def aclose(self):
        pass


async def cycle(seeded, tmp_path, pages, notifier=None):
    return await run_cycle(settings(tmp_path), seeded, pages, notifier=notifier)


async def test_first_cycle_is_baseline_second_detects_and_notifies(seeded, tmp_path):
    pages, sent = Fetch(), []

    async def notifier(events):
        sent.append(events)

    r1 = await cycle(seeded, tmp_path, pages, notifier)
    assert r1["events"] == 0 and not r1["notified"] and sent == []
    pages.pages["deposits"] = mutate(DEPOSITS, "От 12,00% до 14,10%", "От 11,00% до 13,50%")
    r2 = await cycle(seeded, tmp_path, pages, notifier)
    assert r2["events"] >= 1 and r2["notified"] and len(sent) == 1
    assert sent[0][0]["severity"] == "rate" and "Мой доход" in sent[0][0]["text"]


async def test_info_only_changes_are_counted_but_not_pushed(seeded, tmp_path):
    pages, sent = Fetch(), []
    await cycle(seeded, tmp_path, pages)
    pages.pages["deposits"] = mutate(DEPOSITS, "Доступен для всех", "Доступен для всех клиентов банка")
    r = await cycle(seeded, tmp_path, pages, AsyncMock(side_effect=lambda e: sent.append(e)))
    assert r["events"] == 1 and sent == [] and not r["notified"]


async def test_notifier_failure_does_not_break_cycle(seeded, tmp_path):
    pages = Fetch()
    await cycle(seeded, tmp_path, pages)
    pages.pages["deposits"] = mutate(DEPOSITS, "От 12,00% до 14,10%", "От 1% до 2%")
    r = await cycle(seeded, tmp_path, pages, AsyncMock(side_effect=RuntimeError("telegram down")))
    assert r["events"] >= 1 and r["notified"] is False


def test_format_events_escapes_and_limits():
    ev = [{"severity": "rate", "text": f"ПСБ · <b>X{i}</b> & Y"} for i in range(30)]
    text = format_events(ev)
    assert "&lt;b&gt;X0&lt;/b&gt; &amp; Y" in text and "и ещё 5" in text and "<b>X0" not in text


def test_telegram_notifier_requires_explicit_opt_in(tmp_path):
    assert telegram_notifier(settings(tmp_path, bot_token="1:x", allowed_telegram_ids="5")) is None  # NOTIFY_CHANGES выключен
    assert telegram_notifier(settings(tmp_path, notify_changes=True, allowed_telegram_ids="5")) is None  # нет токена
    assert telegram_notifier(settings(tmp_path, notify_changes=True, bot_token="1:x")) is None  # нет получателей
    assert telegram_notifier(settings(tmp_path, notify_changes=True, bot_token="1:x", allowed_telegram_ids="5")) is not None


REDIS = "redis://localhost:56379/0"


def _redis_up() -> bool:
    try:
        import redis

        return bool(redis.Redis.from_url(REDIS, socket_connect_timeout=1).ping())
    except Exception:
        return False


@pytest.mark.skipif(not _redis_up(), reason="нет Redis на localhost:56379")
def test_redis_lock_blocks_concurrent_cycle_and_releases():
    with redis_lock(REDIS) as first:
        assert first is True
        with redis_lock(REDIS) as second:
            assert second is False  # параллельный цикл не стартует
    with redis_lock(REDIS) as again:
        assert again is True  # замок освобождён


def test_redis_lock_unavailable_redis_degrades_with_warning(caplog):
    with redis_lock("redis://localhost:1/0") as got:
        assert got is True
    assert "Redis недоступен" in caplog.text


def test_celery_schedule_and_task_skips_when_locked(monkeypatch):
    from contextlib import contextmanager

    import bankintel.tasks as tasks

    assert tasks.app.conf.beat_schedule["monitor-cycle"]["schedule"] == tasks.settings.monitor_interval_hours * 3600

    @contextmanager
    def locked(url):
        yield False

    monkeypatch.setattr(tasks, "redis_lock", locked)
    assert tasks.monitor_cycle.apply().get() == {"skipped": "предыдущий цикл ещё выполняется"}

    @contextmanager
    def free(url):
        yield True

    monkeypatch.setattr(tasks, "redis_lock", free)
    monkeypatch.setattr(tasks, "_cycle", AsyncMock(return_value={"events": 3}))
    assert tasks.monitor_cycle.apply().get() == {"events": 3}


def test_parse_days():
    assert parse_days(None) == 7 and parse_days("30") == 30
    for bad in ("0", "91", "abc", "-1"):
        with pytest.raises(ValueError):
            parse_days(bad)


async def test_changes_command_states(seeded, tmp_path):
    async with seeded() as s:
        assert "ещё не выполнялся" in await changes_text(s, 7)
    pages = Fetch()
    await cycle(seeded, tmp_path, pages)
    async with seeded() as s:
        t = await changes_text(s, 7)
        assert "Изменений не обнаружено" in t and "базовая линия" in t
    pages.pages["deposits"] = mutate(DEPOSITS, "От 12,00% до 14,10%", "От 11,00% до 13,50%")
    await cycle(seeded, tmp_path, pages)
    m = AsyncMock()
    await handlers.cmd_changes(m, CommandObject(command="changes", args="7"), session_factory=seeded)
    text = m.answer.call_args.args[0]
    assert "ПСБ · Мой доход" in text and "«От 11,00% до 13,50%»" in text and "🔴" in text and "ЛНР не подтверждена" in text
    m = AsyncMock()
    await handlers.cmd_changes(m, CommandObject(command="changes", args="999"), session_factory=seeded)
    assert "от 1 до 90" in m.answer.call_args.args[0]


async def test_report_contains_changes_section(seeded, tmp_path):
    pages = Fetch()
    await cycle(seeded, tmp_path, pages)
    pages.pages["deposits"] = mutate(DEPOSITS, "От 12,00% до 14,10%", "От 11,00% до 13,50%")
    await cycle(seeded, tmp_path, pages)
    async with seeded() as s:
        ds = await build_dataset(s, datetime.now(UTC))
    items = ds["changes"]["items"]
    assert items and any("Мой доход" in c["text"] for c in items)
    from bankintel.reports.render_html import render_html
    from bankintel.reports.render_xlsx import render_xlsx
    from openpyxl import load_workbook
    import io

    assert "Изменения за 30 дней" in render_html(ds) and "«От 11,00% до 13,50%»" in render_html(ds)
    ws = load_workbook(io.BytesIO(render_xlsx(ds)))["Изменения"]
    assert any("Мой доход" in str(r[5].value) for r in ws.iter_rows(min_row=2))
