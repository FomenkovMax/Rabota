"""Продукты — из фрагментов реальных страниц ПСБ/ВТБ (24.10.2026), собранных штатным коллектором."""

import io
import json
from datetime import UTC, datetime
from decimal import Decimal as D
from pathlib import Path
from unittest.mock import AsyncMock

import pytest
from sqlalchemy import select

from bankintel.bot import handlers
from bankintel.collector import collect_bank
from bankintel.db.models import Bank, Product
from bankintel.domain.enums import FetchOutcome
from bankintel.formatting import money
from bankintel.reports.dataset import DatasetError, build_dataset, dataset_id, validate_dataset
from bankintel.reports.service import generate_reports
from bankintel.sources.base import FetchDiagnostic, FetchedDocument

FIX = Path(__file__).parent / "fixtures"
NOW = datetime(2026, 10, 24, 12, tzinfo=UTC)


class _Fetcher:
    PAGES = {"psb": {"savings_accounts": "psb_savings_cards.html"}, "vtb": {}}

    async def fetch(self, source):
        if "vtb.ru" in source.url:
            name = "vtb_cards.html"
        else:
            name = "psb_savings_cards.html" if source.key == "savings_accounts" else "psb_deposits_cards.html"
        return FetchedDocument(source.url, (FIX / name).read_bytes(), "text/html"), FetchDiagnostic(source.key, source.url, FetchOutcome.OK, 200)


@pytest.fixture
async def loaded(seeded, tmp_path):
    async with seeded() as s:
        for slug in ("psb", "vtb"):
            await collect_bank(s, await s.scalar(select(Bank).where(Bank.slug == slug)), _Fetcher(), tmp_path / "raw")
    return seeded


async def test_dataset_is_valid_and_traceable(loaded):
    async with loaded() as s:
        ds = await build_dataset(s, NOW)
    m = ds["meta"]
    assert m["banks_with_data"] == 2 and m["banks_total"] == 5 and m["lnr_confirmed"] == 0
    assert len(ds["catalog"]["deposit"]) == 13 and len(ds["catalog"]["savings"]) == 6
    known = {d["document_id"] for d in ds["sources"]}
    assert all(p["document_ids"] and set(p["document_ids"]) <= known for p in ds["catalog"]["deposit"] + ds["catalog"]["savings"])
    sber = next(c for c in ds["coverage"] if c["slug"] == "sber")
    assert sber["status"] == "no_data" and sber["reason"]  # причина указана
    json.dumps(ds, ensure_ascii=False)  # набор JSON-совместим


async def test_dataset_id_depends_on_content_not_time(loaded):
    async with loaded() as s:
        a = await build_dataset(s, NOW)
        b = await build_dataset(s, datetime(2026, 10, 25, tzinfo=UTC))
    assert a["meta"]["dataset_id"] == b["meta"]["dataset_id"] and a["meta"]["generated_at"] != b["meta"]["generated_at"]
    b["catalog"]["deposit"][0]["rate_max"] = "99"
    assert dataset_id(b) != a["meta"]["dataset_id"]


async def test_empty_database_refuses_report(seeded, tmp_path):
    async with seeded() as s:
        with pytest.raises(DatasetError, match="нет ни одного продукта"):
            await generate_reports(s, tmp_path)
    assert list(tmp_path.iterdir()) == []


async def test_validation_rejects_untraceable_and_synthetic(loaded):
    async with loaded() as s:
        ds = await build_dataset(s, NOW)
    bad = json.loads(json.dumps(ds))
    bad["catalog"]["deposit"][0]["document_ids"] = []
    with pytest.raises(DatasetError, match="нет ссылки на сохранённый документ"):
        validate_dataset(bad)
    bad = json.loads(json.dumps(ds))
    bad["catalog"]["deposit"][0]["name"] = "SYNTHETIC-PRODUCT"
    with pytest.raises(DatasetError, match="тестовые данные"):
        validate_dataset(bad)
    bad = json.loads(json.dumps(ds))
    bad["catalog"]["deposit"][0]["rate_max"] = "250"
    with pytest.raises(DatasetError, match="вне 0–100"):
        validate_dataset(bad)
    bad = json.loads(json.dumps(ds))
    bad["scenarios"][0]["open"][0]["income_min"] = "999999999"
    with pytest.raises(DatasetError, match="интервал дохода"):
        validate_dataset(bad)


async def test_all_formats_built_from_same_data(loaded, tmp_path):
    fitz = pytest.importorskip("fitz")
    from openpyxl import load_workbook
    from pptx import Presentation

    async with loaded() as s:
        bundle = await generate_reports(s, tmp_path / "out", NOW)
    ds, did = bundle.dataset, bundle.dataset_id
    assert set(bundle.paths) == {"pdf", "html", "pptx", "xlsx"} and all(p.stat().st_size > 5000 for p in bundle.paths.values())

    sc = next(x for x in ds["scenarios"] if x["months"] == 12 and D(x["amount"]) == 500000)
    leader = sc["open"][0]
    leader_income = money(leader["income_max"])
    assert leader["name"] and D(leader["income_max"]) == D("75000.00")  # 500 000 × 15%, независимая ручная проверка

    html = bundle.paths["html"].read_text(encoding="utf-8")
    import re

    from bankintel.reports import charts
    from bankintel.reports.render_html import _render

    pdf_source = _render(ds, "pdf", lambda rows, uid: charts.svg_range_chart(rows, uid=uid))
    ids = re.findall(r'<pattern id="([^"]+)"', pdf_source)
    assert len(ids) >= 5 and len(ids) == len(set(ids))  # у каждого графика свои SVG-id
    assert did in html and leader_income in html and "Демо-прототип" in html and "plotly" in html.lower()
    assert "Сбербанк" in html and "блокир" in html.lower() or "недоступен автоматическому сбору" in html

    pdf = "".join(p.get_text() for p in fitz.open(bundle.paths["pdf"]))
    assert did in pdf and leader_income in pdf and "не подтверждена" in pdf

    prs = Presentation(bundle.paths["pptx"])
    texts, cats = [], []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                texts.append(sh.text_frame.text)
            if getattr(sh, "has_chart", False) and sh.has_chart:
                cats += list(sh.chart.plots[0].categories)
    pptx_text = "\n".join(texts)
    assert did in pptx_text and "Сбербанк" in pptx_text
    assert any(leader["name"] in c and money(leader["income_max"]) in c for c in cats) or leader_income in pptx_text
    assert len(prs.slides) >= 9

    wb = load_workbook(bundle.paths["xlsx"])
    assert did in wb.properties.description
    xl_rows = [[c.value for c in r] for r in wb["Сценарии"].iter_rows(min_row=2)]
    xl_keys = {(r[0], r[6], r[7], r[13]) for r in xl_rows}  # сценарий, банк, продукт, доход макс. (число)
    ds_keys = {(s["label"], r["bank"], r["name"], float(D(r["income_max"]))) for s in ds["scenarios"] for r in s["open"] + s["restricted"]}
    assert xl_keys == ds_keys  # XLSX содержит ровно те же строки, что набор данных
    assert all(isinstance(r[13], (int, float)) for r in xl_rows)  # числа, а не строки


async def test_report_handler_sends_four_documents(loaded):
    m = AsyncMock()
    m.answer_document = AsyncMock()
    await handlers.cmd_report(m, session_factory=loaded)
    assert m.answer_document.await_count == 4
    captions = [c.kwargs["caption"] for c in m.answer_document.await_args_list]
    assert {c.split(" · ")[0] for c in captions} == {"PDF", "HTML", "PPTX", "XLSX"}
    assert all("не подтверждена" in c for c in captions)
    assert len({c.split("Набор данных ")[1].split(".")[0] for c in captions}) == 1  # один набор данных во всех файлах


async def test_report_handler_explains_when_no_data(seeded):
    m = AsyncMock()
    m.answer_document = AsyncMock()
    await handlers.cmd_report(m, session_factory=seeded)
    assert m.answer_document.await_count == 0
    assert "Отчёт не сформирован" in m.answer.call_args.args[0] and "нет ни одного продукта" in m.answer.call_args.args[0]


async def test_stale_data_flagged_when_last_fetch_failed(loaded):
    from bankintel.db.models import FetchAttempt, Source

    async with loaded() as s:
        src = await s.scalar(select(Source).join(Bank, Bank.id == Source.bank_id).where(Bank.slug == "psb", Source.key == "deposits"))
        s.add(FetchAttempt(source_id=src.id, outcome="tls_error", error="cert"))
        await s.commit()
        ds = await build_dataset(s, NOW)
    psb = next(c for c in ds["coverage"] if c["slug"] == "psb")
    assert psb["status"] == "data" and "устаревшими" in psb["reason"] and "tls_error" in psb["reason"]
    vtb = next(c for c in ds["coverage"] if c["slug"] == "vtb")
    assert vtb["reason"] is None
