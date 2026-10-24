"""PPTX: нативные (редактируемые) диаграммы и таблицы PowerPoint из того же набора данных."""

import io
from decimal import Decimal

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_PATTERN
from pptx.util import Emu, Inches, Pt

from bankintel.formatting import money, pct
from bankintel.reports import charts
from bankintel.reports.dataset import fmt_date, fmt_dt

INK, INK2, MUTED = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E), RGBColor(0x6D, 0x6C, 0x68)
FONT = "DejaVu Sans"
W, H = Inches(13.333), Inches(7.5)
LNR = {"confirmed": "подтверждена", "unavailable": "недоступен", "not_confirmed": "не подтверждена", "unknown": "неизвестно"}


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.lstrip("#").upper())


def _text(slide, x, y, w, h, text, size=14, bold=False, color=INK, wrap=True):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
        r.font.color.rgb = color
        p.space_after = Pt(4)
    return box


def _slide(prs, ds, title, subtitle=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _text(s, Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.7), title, 24, True)
    if subtitle:
        _text(s, Inches(0.5), Inches(0.95), Inches(12.3), Inches(0.5), subtitle, 12, False, MUTED)
    _text(s, Inches(0.5), Inches(7.05), Inches(12.3), Inches(0.35),
          f"Banking Intelligence AI · набор данных {ds['meta']['dataset_id']} · демо-прототип · доступность в ЛНР не подтверждена · без данных Сбербанка", 9, False, MUTED)
    return s


def _table(slide, x, y, w, rows, header, col_w, size=10, row_h=0.3):
    shape = slide.shapes.add_table(len(rows) + 1, len(header), x, y, w, Inches(row_h * (len(rows) + 1)))
    t = shape.table
    for i, cw in enumerate(col_w):
        t.columns[i].width = Emu(int(w * cw / sum(col_w)))
    for j, h in enumerate(header):
        c = t.cell(0, j)
        c.text = h
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            t.cell(i, j).text = str(v)
    for r in t.rows:
        r.height = Inches(row_h)
        for c in r.cells:
            c.margin_top = c.margin_bottom = Inches(0.02)
            for p in c.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size, run.font.name = Pt(size), FONT
    for i, r in enumerate(t.rows):  # нейтральные заливки вместо синей темы по умолчанию
        for c in r.cells:
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor(0xEE, 0xEE, 0xEA) if i == 0 else RGBColor(0xFF, 0xFF, 0xFF)
            for p in c.text_frame.paragraphs:
                for run in p.runs:
                    run.font.color.rgb, run.font.bold = INK, i == 0
    return t


def _scenario_chart(slide, rows, x, y, w, h):
    top = max(r["hi"] for r in rows)
    eps = top * Decimal("0.012")
    cats, base, span = [], [], []
    for r in reversed(rows):  # бары рисуются снизу вверх
        cats.append(f"{r['label']}  ({charts.value_label(r)})")
        if r["lo"] is not None and r["hi"] > r["lo"]:
            base.append(float(r["lo"])); span.append(float(r["hi"] - r["lo"]))
        else:  # нижняя граница не указана / нулевой интервал: тонкая отметка у верхней границы
            base.append(float(r["hi"] - eps)); span.append(float(eps))
    data = CategoryChartData()
    data.categories = cats
    data.add_series("Нижняя граница (скрыта)", base)
    data.add_series("Интервал дохода, ₽", span)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, x, y, w, h, data)
    ch = gf.chart
    ch.has_legend = False
    ch.font.size, ch.font.name = Pt(10), FONT
    ch.plots[0].gap_width = 60
    ch.series[0].format.fill.background()
    ch.series[0].format.line.fill.background()
    ordered = list(reversed(rows))
    for i, r in enumerate(ordered):
        fmt = ch.series[1].points[i].format
        color = _rgb(charts.bank_color(r["slug"]))
        if r["restricted"]:
            fmt.fill.patterned()
            fmt.fill.pattern = MSO_PATTERN.WIDE_UPWARD_DIAGONAL
            fmt.fill.fore_color.rgb, fmt.fill.back_color.rgb = color, RGBColor(0xFC, 0xFC, 0xFB)
            fmt.line.color.rgb = color
        else:
            fmt.fill.solid()
            fmt.fill.fore_color.rgb = color
    va = ch.value_axis
    va.minimum_scale = 0
    axis_max = charts._nice_max(top)
    va.maximum_scale, va.major_unit = float(axis_max), float(axis_max / 4)
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xE3, 0xE2, 0xDD)
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = '# ##0 "₽"', False
    return ch


def render_pptx(ds: dict) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    m = ds["meta"]

    s = prs.slides.add_slide(prs.slide_layouts[6])
    _text(s, Inches(0.7), Inches(2.2), Inches(12), Inches(1.4), "Конкурентный анализ вкладов и накопительных счетов — ЛНР", 34, True)
    _text(s, Inches(0.7), Inches(3.9), Inches(12), Inches(1.5), [
        f"Данные от {fmt_date(m['data_from'])} · сформировано {fmt_dt(m['generated_at'])}",
        f"Набор данных {m['dataset_id']} · версия {m['app_version']}",
        "ДЕМО-ПРОТОТИП: охвачены ПСБ и ВТБ; данных Сбербанка (эталона) нет; доступность в ЛНР не подтверждена.",
    ], 14, False, INK2)

    s = _slide(prs, ds, "Главное", "Верхние границы дохода по данным сайтов; ставка для клиента зависит от условий")
    kp = [(f"{m['banks_with_data']} из {m['banks_total']}", "банков с данными"),
          (f"{len(ds['catalog']['deposit'])} + {len(ds['catalog']['savings'])}", "вкладов + счетов"),
          (f"{m['lnr_confirmed']} из {m['products_total']}", "доступность в ЛНР подтверждена"),
          (str(len(ds["sources"])), "документов-источников")]
    for i, (v, l) in enumerate(kp):
        _text(s, Inches(0.5 + i * 3.1), Inches(1.5), Inches(3), Inches(0.7), v, 30, True)
        _text(s, Inches(0.5 + i * 3.1), Inches(2.2), Inches(3), Inches(0.4), l, 11, False, INK2)
    _text(s, Inches(0.5), Inches(3.0), Inches(12.3), Inches(3.8), ["• " + t for t in ds["insights"]], 14)

    s = _slide(prs, ds, "Охват данных")
    rows = [[c["name"] + (" (эталон)" if c["is_benchmark"] else ""), "есть данные" if c["status"] == "data" else "нет данных",
             c["deposits"], c["savings"], c["reason"] or (f"загрузка {fmt_dt(c['last_fetch']['at'])}: {c['last_fetch']['outcome']}" if c["last_fetch"] else "")]
            for c in ds["coverage"]]
    _table(s, Inches(0.5), Inches(1.3), Inches(12.3), rows, ["Банк", "Статус", "Вклады", "Счета", "Причина / последняя загрузка"], [2, 1.4, 0.8, 0.8, 7], 11, 0.5)

    for sc in ds["scenarios"]:
        rows = charts.chart_rows(sc)
        s = _slide(prs, ds, f"Сценарий: {sc['label']}", "Интервал дохода за срок: простой процент по ставке сайта, до налогов; штриховка — особые условия; тонкая метка — нижняя граница не указана")
        if rows:
            _scenario_chart(s, rows, Inches(0.4), Inches(1.45), Inches(8.6), Inches(5.4))
            lines = [f"{it['bank']}" for it in charts.legend_items(rows)]
            _text(s, Inches(9.2), Inches(1.5), Inches(3.8), Inches(0.4), "Цвет = банк: " + ", ".join(lines), 11, False, INK2)
            if sc["gap"]:
                g = sc["gap"]
                _text(s, Inches(9.2), Inches(2.3), Inches(3.8), Inches(1.6),
                      f"Среди предложений для всех: {g['leader']} опережает {g['second']} на {money(g['amount'])} по верхней границе.", 13, True)
            _text(s, Inches(9.2), Inches(4.0), Inches(3.8), Inches(2.9),
                  [f"Исключено из сценария: {len(sc['excluded'])}"] + ([f"Нет данных: {', '.join(sc['banks_without_data'])}"] if sc["banks_without_data"] else []), 11, False, INK2)
        else:
            _text(s, Inches(0.5), Inches(2), Inches(12), Inches(1), "Подходящих продуктов в этом сценарии нет.", 16)

    mon = ds["changes"]["monitoring"]
    s = _slide(prs, ds, f"Изменения за {ds['changes']['window_days']} дней",
               f"Наблюдение с {fmt_dt(mon['since'])}; сравнений с предыдущим запуском: {mon['comparisons']}" + ("; пока базовая линия (первый сбор)" if not mon["comparisons"] else ""))
    items = ds["changes"]["items"][:10]
    if items:
        _table(s, Inches(0.5), Inches(1.4), Inches(12.3), [[fmt_dt(c["detected_at"]), c["text"]] for c in items], ["Когда", "Что изменилось"], [1.8, 10.5], 10, 0.45)
    else:
        _text(s, Inches(0.5), Inches(1.6), Inches(12), Inches(1), "Изменений за период не обнаружено.", 16)

    def prod_rows(items, deposit):
        out = []
        for p in items:
            amt = (f"от {money(p['amount_min'])}" if p["amount_min"] else "") + (f" до {money(p['amount_max'])}" if p["amount_max"] else "")
            tags = ", ".join(t for t, k in (("особые условия", "restricted_audience"), ("плавающая", "floating_rate"), ("ограничено по времени", "limited_time")) if k in p["flags"])
            row = [p["name"], p["rate_text"] or "—"] + ([p["term_text"] or "—"] if deposit else []) + [amt or "—", tags, LNR[p["lnr_status"]]]
            out.append(row)
        return out

    for bank in dict.fromkeys(p["bank"] for p in ds["catalog"]["deposit"]):
        items = [p for p in ds["catalog"]["deposit"] if p["bank"] == bank]
        s = _slide(prs, ds, f"Вклады: {bank}", "Ставки и условия — как опубликовано на сайтах; не гарантия для конкретной суммы и срока")
        _table(s, Inches(0.5), Inches(1.4), Inches(12.3), prod_rows(items, True)[:14], ["Продукт", "Ставка (сайт)", "Срок", "Сумма", "Пометки", "ЛНР"], [2.4, 1.9, 2.3, 2.5, 2.1, 1.9], 9, 0.36)
    for bank in dict.fromkeys(p["bank"] for p in ds["catalog"]["savings"]):
        items = [p for p in ds["catalog"]["savings"] if p["bank"] == bank]
        s = _slide(prs, ds, f"Накопительные счета: {bank}", "«До X%» — максимум при выполнении условий банка; доход по счетам не рассчитывается")
        _table(s, Inches(0.5), Inches(1.4), Inches(12.3), prod_rows(items, False), ["Продукт", "Ставка (сайт)", "Сумма", "Пометки", "ЛНР"], [2.8, 2.3, 2.8, 2.4, 2.0], 10, 0.4)

    s = _slide(prs, ds, "Методика и ограничения")
    _text(s, Inches(0.5), Inches(1.2), Inches(6.1), Inches(5.7), ["Методика"] + ["• " + t for t in ds["methodology"]], 11)
    _text(s, Inches(6.8), Inches(1.2), Inches(6.1), Inches(5.7), ["Ограничения"] + ["• " + t for t in ds["limitations"]], 11)

    buf = io.BytesIO()
    prs.core_properties.title, prs.core_properties.comments = "Банковская конкуренция, ЛНР", f"dataset_id={m['dataset_id']}"
    prs.save(buf)
    return buf.getvalue()
