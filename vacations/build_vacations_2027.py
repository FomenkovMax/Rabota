import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment

OUT = "/home/user/Rabota/vacations/График_отпусков_2027_ЛНР.xlsx"
YEAR = 2027
N = 40               # строк под сотрудников
FIRST = 5            # первая строка данных
LAST = FIRST + N - 1
PARTS = 4

F = "Arial"
def font(**kw): return Font(name=F, **kw)
def fill(c): return PatternFill("solid", start_color=c, end_color=c)
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

GREEN_DARK = "1F7A3E"   # фирменный зелёный
GREEN_LIGHT = "DFF2E4"
INPUT = "FFF7D6"        # ячейки для ввода
GRAY = "F2F2F2"
HOL = "F8D7DA"
RED = "E5534B"
OK = "C6EFCE"
BAD = "FFC7CE"

wb = Workbook()

# ---------------- Праздники ----------------
hs = wb.active
hs.title = "Праздники"
holidays = [
    (dt.date(YEAR, 1, d), "Новогодние каникулы", "ст. 112 ТК РФ") for d in (1, 2, 3, 4, 5, 6, 8)
] + [
    (dt.date(YEAR, 1, 7), "Рождество Христово", "ст. 112 ТК РФ"),
    (dt.date(YEAR, 2, 23), "День защитника Отечества", "ст. 112 ТК РФ"),
    (dt.date(YEAR, 3, 8), "Международный женский день", "ст. 112 ТК РФ"),
    (dt.date(YEAR, 5, 1), "Праздник Весны и Труда", "ст. 112 ТК РФ"),
    (dt.date(YEAR, 5, 9), "День Победы", "ст. 112 ТК РФ"),
    (dt.date(YEAR, 5, 12), "День Луганской Народной Республики", "Указ Главы ЛНР (янв. 2023), региональный"),
    (dt.date(YEAR, 6, 12), "День России", "ст. 112 ТК РФ"),
    (dt.date(YEAR, 9, 30), "День воссоединения ЛНР с Россией", "Указ Главы ЛНР (янв. 2023), региональный"),
    (dt.date(YEAR, 11, 4), "День народного единства", "ст. 112 ТК РФ"),
]
holidays.sort()
hs["A1"], hs["B1"], hs["C1"] = "Дата", "Праздник", "Основание"
for c in "ABC":
    hs[f"{c}1"].font = font(bold=True, color="FFFFFF")
    hs[f"{c}1"].fill = fill(GREEN_DARK)
    hs[f"{c}1"].alignment = CENTER
for i, (d, name, src) in enumerate(holidays, start=2):
    hs.cell(i, 1, d).number_format = "DD.MM.YYYY"
    hs.cell(i, 2, name)
    hs.cell(i, 3, src)
HOL_LAST = 41
for r in range(2, HOL_LAST + 1):
    for c in range(1, 4):
        cell = hs.cell(r, c)
        cell.font = font()
        cell.border = BOX
        if c == 1:
            cell.fill = fill(INPUT)
            cell.number_format = "DD.MM.YYYY"
hs.column_dimensions["A"].width = 13
hs.column_dimensions["B"].width = 40
hs.column_dimensions["C"].width = 44
hs["E1"] = "Нерабочие праздничные дни, попавшие в отпуск, в число дней отпуска не входят — отпуск продлевается (ст. 120 ТК РФ)."
hs["E2"] = "Переносы выходных (постановление Правительства) на подсчёт не влияют — отпуск считается в календарных днях."
hs["E3"] = "Список можно править: добавьте/удалите даты в столбце A (до строки 41) — всё пересчитается автоматически."
hs["E4"] = "Региональные дни ЛНР (12.05 и 30.09) — по указу Главы ЛНР от января 2023 г. Перед утверждением графика сверьте с кадрами."
for r in range(1, 5):
    hs[f"E{r}"].font = font(italic=True, color="595959")
hs.freeze_panes = "A2"
wb.defined_names["HOLIDAYS"] = DefinedName("HOLIDAYS", attr_text=f"'Праздники'!$A$2:$A${HOL_LAST}")

# ---------------- Отпуска ----------------
ws = wb.create_sheet("Отпуска", 0)
ws.sheet_view.showGridLines = False
NCOL_PART0 = 4
col_total = NCOL_PART0 + 4 * PARTS          # 20
col_max = col_total + 1
col_ovl = col_total + 2
col_status = col_total + 3

ws["A1"] = f"График отпусков на {YEAR} год — ЛНР"
ws["A1"].font = font(bold=True, size=16, color=GREEN_DARK)
ws["A2"] = ("Как заполнять: жёлтые ячейки — ФИО, подразделение, дата начала и кол-во календарных дней каждой части. "
            "Дата окончания считается сама: праздники ЛНР/РФ внутри отпуска не учитываются и отпуск продлевается. "
            "Правила: всего ≥ 28 дн., одна из частей ≥ 14 дн. Строка 5 — пример, её можно перезаписать. "
            "Пересечения — на листе «Календарь».")
ws["A2"].font = font(size=9, color="595959")
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_status)
ws.row_dimensions[2].height = 30

def hdr(cell, text):
    cell.value = text
    cell.font = font(bold=True, color="FFFFFF", size=10)
    cell.fill = fill(GREEN_DARK)
    cell.alignment = CENTER
    cell.border = BOX

for c, t in [(1, "№"), (2, "ФИО сотрудника"), (3, "Подразделение / должность"),
             (col_total, "Итого дней"), (col_max, "Макс. часть"),
             (col_ovl, "Дней пересечений с коллегами"), (col_status, "Проверка")]:
    hdr(ws.cell(3, c), t)
    ws.merge_cells(start_row=3, start_column=c, end_row=4, end_column=c)
    ws.cell(4, c).border = BOX
for p in range(PARTS):
    c0 = NCOL_PART0 + 4 * p
    hdr(ws.cell(3, c0), f"Часть {p + 1}")
    ws.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c0 + 3)
    for j, t in enumerate(["Начало", "Дней", "Окончание", "Праздн. внутри"]):
        hdr(ws.cell(4, c0 + j), t)
        ws.cell(4, c0 + j).fill = fill("2E9A52")
ws.row_dimensions[3].height = 22
ws.row_dimensions[4].height = 30

for r in range(FIRST, LAST + 1):
    ws.cell(r, 1, r - FIRST + 1)
    for p in range(PARTS):
        c0 = NCOL_PART0 + 4 * p
        S, D, E = (f"{L(c0)}{r}", f"{L(c0 + 1)}{r}", f"{L(c0 + 2)}{r}")
        ws.cell(r, c0 + 2, f'=IF(OR({S}="",{D}=""),"",{S}+{D}-1+SUMPRODUCT((HOLIDAYS>={S})*(HOLIDAYS-COUNTIFS(HOLIDAYS,">="&{S},HOLIDAYS,"<"&HOLIDAYS)<={S}+{D}-1)))')
        ws.cell(r, c0 + 3, f'=IF({E}="","",{E}-{S}+1-{D})')
    dcols = [f"{L(NCOL_PART0 + 4 * p + 1)}{r}" for p in range(PARTS)]
    tot = f"{L(col_total)}{r}"
    mx = f"{L(col_max)}{r}"
    ws.cell(r, col_total, f"=IF(COUNT({','.join(dcols)})=0,\"\",SUM({','.join(dcols)}))")
    ws.cell(r, col_max, f"=IF({tot}=\"\",\"\",MAX({','.join(dcols)}))")
    first_d, last_d = L(3), L(3 + 364)
    ws.cell(r, col_ovl,
            f"=IF({tot}=\"\",\"\",SUMPRODUCT((Календарь!${first_d}{r}:${last_d}{r}=1)*(Календарь!${first_d}${LAST + 1}:${last_d}${LAST + 1}>=2)))")
    ws.cell(r, col_status,
            f"=IF(AND($B{r}=\"\",{tot}=\"\"),\"\",IF({tot}=\"\",\"Нет дат\",IF({tot}<28,\"Меньше 28 дн.\","
            f"IF({mx}<14,\"Нет части ≥14 дн.\",\"✔ OK\"))))")
    for c in range(1, col_status + 1):
        cell = ws.cell(r, c)
        cell.font = font(size=10)
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center" if c not in (2, 3) else "left", vertical="center")
        rel = c - NCOL_PART0
        is_input = c in (2, 3) or (0 <= rel < 4 * PARTS and rel % 4 in (0, 1))
        if is_input:
            cell.fill = fill(INPUT)
        elif 0 <= rel < 4 * PARTS:
            cell.fill = fill(GRAY)
        if 0 <= rel < 4 * PARTS and rel % 4 in (0, 2):
            cell.number_format = "DD.MM.YYYY"
    ws.cell(r, col_total).font = font(size=10, bold=True)

# пример
ws.cell(FIRST, 2, "Иванов Иван Иванович (пример)")
ws.cell(FIRST, 3, "УСИП, эксперт")
ws.cell(FIRST, 4, dt.date(YEAR, 5, 4)); ws.cell(FIRST, 5, 14)
ws.cell(FIRST, 8, dt.date(YEAR, 8, 16)); ws.cell(FIRST, 9, 14)
ws.cell(FIRST, 2).comment = Comment("Пример: часть 1 захватывает 9 и 12 мая — окончание сдвинулось на 2 дня.", "Шаблон")

# ширины
ws.column_dimensions["A"].width = 4
ws.column_dimensions["B"].width = 32
ws.column_dimensions["C"].width = 22
for p in range(PARTS):
    c0 = NCOL_PART0 + 4 * p
    for j, w in enumerate([11.5, 6, 11.5, 8]):
        ws.column_dimensions[L(c0 + j)].width = w
for c, w in [(col_total, 8), (col_max, 8), (col_ovl, 12), (col_status, 18)]:
    ws.column_dimensions[L(c)].width = w
ws.freeze_panes = ws.cell(FIRST, 4)

# проверки ввода
dv_date = DataValidation(type="date", operator="between",
                         formula1=f"DATE({YEAR},1,1)", formula2=f"DATE({YEAR},12,31)",
                         showErrorMessage=True, errorTitle="Дата",
                         error=f"Введите дату {YEAR} года в формате ДД.ММ.ГГГГ",
                         showInputMessage=True, promptTitle="Начало отпуска", prompt="ДД.ММ.ГГГГ")
dv_days = DataValidation(type="whole", operator="between", formula1="1", formula2="60",
                         showErrorMessage=True, errorTitle="Дни", error="Целое число календарных дней (1–60)")
ws.add_data_validation(dv_date); ws.add_data_validation(dv_days)
for p in range(PARTS):
    c0 = NCOL_PART0 + 4 * p
    dv_date.add(f"{L(c0)}{FIRST}:{L(c0)}{LAST}")
    dv_days.add(f"{L(c0 + 1)}{FIRST}:{L(c0 + 1)}{LAST}")

# условное форматирование
st = f"{L(col_status)}{FIRST}:{L(col_status)}{LAST}"
ws.conditional_formatting.add(st, FormulaRule(formula=[f'${L(col_status)}{FIRST}="✔ OK"'],
                                              fill=fill(OK), font=Font(name=F, color="006100", bold=True)))
ws.conditional_formatting.add(st, FormulaRule(formula=[f'AND(${L(col_status)}{FIRST}<>"",${L(col_status)}{FIRST}<>"✔ OK")'],
                                              fill=fill(BAD), font=Font(name=F, color="9C0006", bold=True)))
ov = f"{L(col_ovl)}{FIRST}:{L(col_ovl)}{LAST}"
ws.conditional_formatting.add(ov, FormulaRule(formula=[f"AND(ISNUMBER({L(col_ovl)}{FIRST}),{L(col_ovl)}{FIRST}>0)"],
                                              fill=fill("FCE4D6"), font=Font(name=F, color="C00000", bold=True)))
for p in range(PARTS):
    c0 = NCOL_PART0 + 4 * p
    rng = f"{L(c0 + 3)}{FIRST}:{L(c0 + 3)}{LAST}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({L(c0 + 3)}{FIRST}),{L(c0 + 3)}{FIRST}>0)"],
                                                   font=Font(name=F, color="C00000", bold=True)))
    # начало выпало на праздник — предупреждение
    rng = f"{L(c0)}{FIRST}:{L(c0)}{LAST}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND({L(c0)}{FIRST}<>\"\",COUNTIF(HOLIDAYS,{L(c0)}{FIRST})>0)"],
                                                   fill=fill("F4B183")))

# ---------------- Календарь ----------------
cs = wb.create_sheet("Календарь", 1)
cs.sheet_view.showGridLines = False
C0 = 3
CL = C0 + 364
cs["A1"] = f"Календарь отпусков {YEAR}"
cs["A1"].font = font(bold=True, size=14, color=GREEN_DARK)
cs["A2"] = "Зелёный — отпуск, красный — пересечение (в отпуске ≥2 чел.), розовый — праздник, серый — выходной."
cs["A2"].font = font(size=9, color="595959")
cs["A2"].alignment = Alignment(wrap_text=True, vertical="top")
cs.merge_cells("A2:B4")
cs.column_dimensions["A"].width = 4
cs.column_dimensions["B"].width = 30

months = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август",
          "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]
d = dt.date(YEAR, 1, 1)
col = C0
while d.year == YEAR:
    if d.day == 1:
        start = col
    cs.column_dimensions[L(col)].width = 3.3
    nxt = d + dt.timedelta(days=1)
    if nxt.month != d.month:
        c = cs.cell(2, start, months[d.month - 1])
        cs.merge_cells(start_row=2, start_column=start, end_row=2, end_column=col)
        c.font = font(bold=True, color="FFFFFF", size=10)
        c.fill = fill(GREEN_DARK if d.month % 2 else "2E9A52")
        c.alignment = CENTER
    d = nxt
    col += 1

cs.cell(3, C0, dt.date(YEAR, 1, 1))
for c in range(C0, CL + 1):
    if c > C0:
        cs.cell(3, c, f"={L(c - 1)}3+1")
    cs.cell(4, c, f"={L(c)}3")
    cs.cell(3, c).number_format = "D"
    cs.cell(4, c).number_format = "[$-419]DDD"
    for r in (3, 4):
        cs.cell(r, c).font = font(size=7, bold=(r == 3))
        cs.cell(r, c).alignment = CENTER
        cs.cell(r, c).border = BOX

for r in range(FIRST, LAST + 1):
    cs.cell(r, 1, f"=Отпуска!A{r}").font = font(size=9)
    cs.cell(r, 2, f'=IF(Отпуска!B{r}="","",Отпуска!B{r})').font = font(size=9)
    cs.cell(r, 1).border = BOX; cs.cell(r, 2).border = BOX
    conds = []
    for p in range(PARTS):
        c0 = NCOL_PART0 + 4 * p
        conds.append(f"AND(X$3>=Отпуска!${L(c0)}{r},X$3<=N(Отпуска!${L(c0 + 2)}{r}))")
    tmpl = "=IF(OR(" + ",".join(conds) + "),1,\"\")"
    for c in range(C0, CL + 1):
        cell = cs.cell(r, c, tmpl.replace("X$3", f"{L(c)}$3"))
        cell.number_format = ";;;"
        cell.border = BOX
CNT = LAST + 1
cs.cell(CNT, 2, "В отпуске, чел.").font = font(bold=True, size=9)
cs.cell(CNT, 2).alignment = Alignment(horizontal="right")
for c in range(C0, CL + 1):
    cell = cs.cell(CNT, c, f"=IF(COUNT({L(c)}{FIRST}:{L(c)}{LAST})=0,\"\",COUNT({L(c)}{FIRST}:{L(c)}{LAST}))")
    cell.font = font(size=7, bold=True)
    cell.alignment = CENTER
    cell.border = BOX

body = f"{L(C0)}{FIRST}:{L(CL)}{LAST}"
fc = L(C0)
cs.conditional_formatting.add(body, FormulaRule(formula=[f"AND({fc}{FIRST}=1,{fc}${CNT}>=2)"], fill=fill(RED), stopIfTrue=True))
cs.conditional_formatting.add(body, FormulaRule(formula=[f"{fc}{FIRST}=1"], fill=fill("3BAA5C"), stopIfTrue=True))
for rng, r0 in ((body, FIRST), (f"{fc}3:{L(CL)}4", 3)):
    cs.conditional_formatting.add(rng, FormulaRule(formula=[f"COUNTIF(HOLIDAYS,{fc}$3)>0"], fill=fill(HOL),
                                                   font=Font(name=F, color="C00000"), stopIfTrue=True))
    cs.conditional_formatting.add(rng, FormulaRule(formula=[f"WEEKDAY({fc}$3,2)>5"], fill=fill("E7E6E6")))
cs.conditional_formatting.add(f"{fc}{CNT}:{L(CL)}{CNT}",
                              FormulaRule(formula=[f"AND(ISNUMBER({fc}{CNT}),{fc}{CNT}>=2)"],
                                          fill=fill(RED), font=Font(name=F, color="FFFFFF", bold=True)))
cs.freeze_panes = cs.cell(FIRST, C0)

ws.sheet_properties.tabColor = GREEN_DARK
cs.sheet_properties.tabColor = "3BAA5C"
hs.sheet_properties.tabColor = "C00000"
ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToHeight = 0

import os
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(OUT)
